from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
import subprocess
import json
import logging
import traceback
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.responses import StreamingResponse
import os
from docx import Document
from io import BytesIO
# from PyPDF2 import PdfReader
import pdfplumber
from dotenv import load_dotenv
from Searching_Top_CVs.vector_bm25_search_revised import initialize_chromadb_client
from Searching_Top_CVs.Process_Upload_Functions.azure_storage_utils import (
    get_blob_container_client,
    download_blob_to_stream
)

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# from search_resume_md11 import search_resumes
from Searching_Top_CVs.vector_bm25_search_revised import process_resumes
# from Searching_Top_CVs.vector_bm25_resume import process_resumes


load_dotenv()

processed_folder = os.getenv('PROCESSED_DATA_FOLDER')

@app.get("/")
async def root():
    return {"message": "Welcome to the application!"}

@app.get("/run-script")
async def run_script():
    result = subprocess.run(["python", "app/script.py"], capture_output=True, text=True)
    if result.returncode == 0:
        return {"output": result.stdout.strip()}
    else:
        return {"error": result.stderr, "code": result.returncode}

class BooleanNameValue(BaseModel):
    BOOLEANNAME: str
    BOOLEANVALUE: str

class BooleanQuery(BaseModel):
    query: str
    n_results: int
    poolValue: int
    toggleJobTitle: bool
    booleanNameValue: List[BooleanNameValue]
    experienceFromValue: int
    experienceToValue: int
    jobQuery: str

@app.post('/search-resumes')
async def search_resumes_endpoint(boolean_query: BooleanQuery):
    try:
        logging.debug(f"Received boolean query: {boolean_query.query} with n_results: {boolean_query.n_results}")
        # results = search_resumes(boolean_query.query, boolean_query.n_results)
        lambda_factor=0.05
        keywords_with_weights = [
            (kw.strip().split(',')[0].strip('" '), float(kw.strip().split(',')[1].strip('" ')) / 100)
            for kw in boolean_query.query.split('AND')
        ]

        experience_from = None if boolean_query.experienceFromValue == 0 else boolean_query.experienceFromValue
        experience_to = None if boolean_query.experienceToValue == 0 else boolean_query.experienceToValue

        boolean_synonym = {}
        for pair in boolean_query.booleanNameValue:
            boolean_synonym[pair.BOOLEANNAME] = [
                value.strip(" '\"") for value in pair.BOOLEANVALUE.split(" OR ")
            ]

        # print(f"Boolean Synonym: {boolean_synonym}")

        collection = await initialize_chromadb_client()

        results = await process_resumes(
            collection,
            boolean_query.toggleJobTitle,
            keywords_with_weights,
            boolean_synonym,
            boolean_query.jobQuery,
            boolean_query.poolValue,
            experience_from,
            experience_to,
            top_n=boolean_query.n_results
        )

        return {"results": results} 
    except Exception as e:
        logging.error(f"Error searching resumes: {e}")
        logging.error(traceback.format_exc())
        raise HTTPException(status_code=500, detail=str(e))

# RAW_DATA_FOLDER = '../chroma/RawData'

@app.get("/list-files")
async def list_files():
    try:
        files = os.listdir(processed_folder)
        return {"files": files}
    except FileNotFoundError as e:
        logging.error(f"Error listing files: {e}")
        raise HTTPException(status_code=404, detail="Directory not found")

# Backup implementation if we need to go back to uploading the image through local folders on the server

# class CORSFileResponse(FileResponse):
#     def __init__(self, *args, **kwargs):
#         super().__init__(*args, **kwargs)
#         self.headers["Access-Control-Allow-Origin"] = "*"
#         self.headers["Access-Control-Allow-Methods"] = "GET, OPTIONS"
#         self.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization"

# @app.get("/view-resume/{filename}")
# async def view_resume(filename: str):
#     file_path = os.path.join(processed_folder, filename)
#     logging.info(f"Attempting to access file: {file_path}")
#     if os.path.exists(file_path):
#         logging.info(f"File found: {file_path}")
#         if filename.endswith('.pdf'):
#             return FileResponse(file_path, media_type="application/pdf", headers={"Content-Disposition": f"inline; filename={filename}"})
#         elif filename.endswith('.docx'):
#             return CORSFileResponse(file_path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"inline; filename={filename}"})
#         else:
#             raise HTTPException(status_code=400, detail="Unsupported file type")
#     else:
#         logging.error(f"File not found: {file_path}")
#         raise HTTPException(status_code=404, detail="File not found")

# @app.get("/view-resume/{filename}")
# async def view_resume(filename: str):
#     """
#     Serve resumes directly from Azure Blob Storage as a stream.
#     """
#     try:
#         # Initialize Azure Blob Storage client
#         connection_string = os.getenv('AZURE_STORAGE_CONNECTION_STRING')
#         container_name = os.getenv('AZURE_BLOB_NAME')
#         container_client = get_blob_container_client(connection_string, container_name)

#         if not container_client:
#             logging.error("Failed to initialize container client.")
#             raise HTTPException(status_code=500, detail="Storage service unavailable")

#         # Construct the blob path
#         blob_name = f"resume/{filename}"

#         # Download the blob as a stream
#         stream = download_blob_to_stream(container_client, blob_name)
#         if not stream:
#             logging.error(f"Blob not found: {blob_name}")
#             raise HTTPException(status_code=404, detail="File not found")

#         # Determine the media type based on the file extension
#         if filename.endswith('.pdf'):
#             media_type = "application/pdf"
#         elif filename.endswith('.docx'):
#             media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
#         else:
#             raise HTTPException(status_code=400, detail="Unsupported file type")

#         # Return the file as a streaming response with CORS headers
#         return StreamingResponse(stream, media_type=media_type, headers={
#             "Content-Disposition": f"inline; filename={filename}",
#             "Access-Control-Allow-Origin": "*",
#             "Access-Control-Allow-Methods": "GET, OPTIONS",
#             "Access-Control-Allow-Headers": "Content-Type, Authorization"
#         })

#     except Exception as e:
#         logging.error(f"Error serving file {filename}: {e}")
#         raise HTTPException(status_code=500, detail="Internal server error")

@app.get("/view-resume/{filename}")
async def view_resume(filename: str):
    """
    Serve resumes directly from Azure Blob Storage as a stream.
    If the file is a PDF, convert it to DOCX before serving.
    """
    try:
        # Initialize Azure Blob Storage client
        connection_string = os.getenv('AZURE_STORAGE_CONNECTION_STRING')
        container_name = os.getenv('AZURE_BLOB_NAME')
        container_client = get_blob_container_client(connection_string, container_name)

        if not container_client:
            logging.error("Failed to initialize container client.")
            raise HTTPException(status_code=500, detail="Storage service unavailable")

        # Construct the blob path
        blob_name = f"resume/{filename}"

        # Download the blob as a stream
        stream = download_blob_to_stream(container_client, blob_name)
        if not stream:
            logging.error(f"Blob not found: {blob_name}")
            raise HTTPException(status_code=404, detail="File not found")

        # Determine the media type based on the file extension
        if filename.endswith('.pdf'):
            # Convert PDF to DOCX
            pdf_stream = BytesIO(stream.read())  # Read the PDF stream into memory
            docx_stream = convert_pdf_to_docx(pdf_stream)

            # Return the DOCX file as a streaming response
            return StreamingResponse(
                docx_stream,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={
                    "Content-Disposition": f"inline; filename={filename.replace('.pdf', '.docx')}",  # Update filename
                    "Access-Control-Allow-Origin": "*",
                    "Access-Control-Allow-Methods": "GET, OPTIONS",
                    "Access-Control-Allow-Headers": "Content-Type, Authorization"
                }
            )
        elif filename.endswith('.docx'):
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")

        # Return the file as a streaming response with CORS headers
        return StreamingResponse(stream, media_type=media_type, headers={
            "Content-Disposition": "inline",  # Ensure the file is displayed in the browser
            "Access-Control-Allow-Origin": "*",
            "Access-Control-Allow-Methods": "GET, OPTIONS",
            "Access-Control-Allow-Headers": "Content-Type, Authorization"
        })

    except Exception as e:
        logging.error(f"Error serving file {filename}: {e}")
        raise HTTPException(status_code=500, detail="Internal server error")
    
def convert_pdf_to_docx(pdf_stream: BytesIO) -> BytesIO:
    """
    Converts a PDF file to a DOCX file in memory.
    """
    document = Document()

    try:
        # Use pdfplumber to extract text from the PDF
        with pdfplumber.open(pdf_stream) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    document.add_paragraph(text)
                else:
                    document.add_paragraph("[Unable to extract text from this page]")

        # Save the DOCX content to a BytesIO stream
        docx_stream = BytesIO()
        document.save(docx_stream)
        docx_stream.seek(0)  # Reset the stream position to the beginning

        return docx_stream

    except Exception as e:
        logging.error(f"Error converting PDF to DOCX: {e}")
        raise HTTPException(status_code=500, detail="Failed to convert PDF to DOCX.")